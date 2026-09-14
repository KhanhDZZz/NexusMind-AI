import os
from pathlib import Path
import pandas as pd
from docx import Document
from pypdf import PdfReader
from PIL import Image

try:
    import pytesseract
except ImportError:
    pytesseract = None

def parse_file(file_path: str) -> str:
    """Tự động nhận diện định dạng tệp và trích xuất nội dung văn bản."""
    ext = Path(file_path).suffix.lower()
    
    # 1. Văn bản thuần, Mã nguồn, SQL
    text_extensions = {
        '.txt', '.md', '.py', '.sql', '.js', '.ts', '.html', '.css',
        '.json', '.xml', '.yaml', '.yml', '.cpp', '.c', '.java', '.cs', '.sh'
    }
    if ext in text_extensions:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    # 2. Tài liệu PDF
    elif ext == '.pdf':
        reader = PdfReader(file_path)
        text = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            text.append(f"--- Trang {i + 1} ---\n{page_text}")
        return "\n\n".join(text)

    # 3. Tài liệu Word (.docx)
    elif ext == '.docx':
        doc = Document(file_path)
        full_text = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                full_text.append(" | ".join([cell.text.strip() for cell in row.cells]))
        return "\n".join(full_text)

    # 4. Bảng tính Excel & CSV (.xlsx, .xls, .csv)
    elif ext in ['.xlsx', '.xls', '.csv']:
        if ext == '.csv':
            df = pd.read_csv(file_path)
            return f"--- Bảng dữ liệu CSV ---\n" + df.to_markdown(index=False)
        else:
            excel_file = pd.ExcelFile(file_path)
            result = []
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                result.append(f"--- Sheet: {sheet_name} ---\n" + df.to_markdown(index=False))
            return "\n\n".join(result)

    # 5. Hình ảnh (OCR với pytesseract)
    elif ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp']:
        if pytesseract:
            try:
                img = Image.open(file_path)
                text = pytesseract.image_to_string(img, lang='vie+eng')
                return f"--- Văn bản trích xuất từ ảnh ({os.path.basename(file_path)}) ---\n" + text
            except Exception as e:
                return f"[Không thể trích xuất chữ từ ảnh: {e}]"
        else:
            return "[Lưu ý: Chưa cấu hình Tesseract OCR để nhận diện chữ trong ảnh]"

    # 6. Tệp âm thanh
    elif ext in ['.mp3', '.wav', '.m4a']:
        return "[Lưu ý: Để đọc tệp âm thanh, vui lòng tích hợp mô hình Speech-to-Text như OpenAI Whisper]"

    return f"[Định dạng tệp {ext} chưa được hỗ trợ trích xuất tự động]"